"""Text extraction for supported product-workbench resources."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _paragraph_text(paragraphs) -> list[str]:
    return [paragraph.text.strip() for paragraph in paragraphs if paragraph.text.strip()]


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return ParsedDocument(text=content.decode("utf-8-sig"))

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return ParsedDocument(
            text="\n\n".join(f"[第 {index} 页]\n{text.strip()}" for index, text in enumerate(pages, 1)),
            page_count=len(pages),
        )

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        blocks = _paragraph_text(document.paragraphs)
        for table_index, table in enumerate(document.tables, 1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            blocks.append(f"[表格 {table_index}]\n" + "\n".join(rows))
        return ParsedDocument(text="\n\n".join(blocks), metadata={"tables": len(document.tables)})

    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(content))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, 1):
            lines: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "").strip()
                if text:
                    lines.append(text)
            if slide.has_notes_slide:
                notes = _paragraph_text(slide.notes_slide.notes_text_frame.paragraphs)
                if notes:
                    lines.append("[备注]\n" + "\n".join(notes))
            slides.append(f"[第 {index} 张幻灯片]\n" + "\n".join(lines))
        return ParsedDocument(text="\n\n".join(slides), page_count=len(slides))

    raise ValueError("仅支持 PPTX、PDF、DOCX、Markdown 和 TXT 文件")


def resource_kind(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return {
        ".pptx": "pptx",
        ".pdf": "pdf",
        ".docx": "docx",
        ".md": "markdown",
        ".markdown": "markdown",
        ".txt": "text",
    }.get(suffix, "text")
