import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

from coursebook_agent.app import app
from coursebook_agent.models import Course, CourseBook, JobState, Lecture, TranscriptSegment
from coursebook_agent.product.models import DatasetCreate, SnapshotCreate
from coursebook_agent.product.parsers import parse_document
from coursebook_agent.product.service import ProductService


class DocumentParserTests(unittest.TestCase):
    def test_docx_and_pptx_extract_structure(self):
        from docx import Document
        from pptx import Presentation

        docx_buffer = BytesIO()
        document = Document()
        document.add_heading("课程讲义", 1)
        document.add_paragraph("假设检验内容")
        document.save(docx_buffer)
        self.assertIn("假设检验内容", parse_document("notes.docx", docx_buffer.getvalue()).text)

        pptx_buffer = BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "统计学"
        slide.placeholders[1].text = "总体与样本"
        presentation.save(pptx_buffer)
        parsed = parse_document("slides.pptx", pptx_buffer.getvalue())
        self.assertEqual(parsed.page_count, 1)
        self.assertIn("总体与样本", parsed.text)

    def test_unsupported_binary_file_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "仅支持"):
            parse_document("legacy.ppt", b"binary")


class ProductServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.service = ProductService(Path(self.temp.name))

    def tearDown(self):
        self.temp.cleanup()

    def test_upload_preview_and_snapshot(self):
        dataset = self.service.create_dataset(DatasetCreate(name="统计学资料"))
        resource = self.service.add_file(dataset.dataset_id, "notes.md", b"# Notes\nHypothesis testing")
        revision = resource.current_revision
        self.assertIsNotNone(revision)
        self.assertEqual(revision.parse_status, "ready")
        self.assertIn("Hypothesis", self.service.preview_revision(revision.revision_id).text)

        snapshot = self.service.create_snapshot(
            dataset.dataset_id,
            SnapshotCreate(resource_revision_ids=[revision.revision_id], label="First run"),
        )
        self.assertEqual(snapshot.resource_count, 1)
        self.assertEqual(snapshot.resource_revision_ids, [revision.revision_id])
        self.assertEqual(snapshot.sha256, self.service.get_snapshot(snapshot.snapshot_id).sha256)

    def test_snapshot_rejects_revision_from_another_dataset(self):
        first = self.service.create_dataset(DatasetCreate(name="First"))
        second = self.service.create_dataset(DatasetCreate(name="Second"))
        revision = self.service.add_file(first.dataset_id, "a.txt", b"a").current_revision
        with self.assertRaisesRegex(ValueError, "不属于"):
            self.service.create_snapshot(second.dataset_id, SnapshotCreate(resource_revision_ids=[revision.revision_id]))

    @patch("coursebook_agent.product.service.ZhiyunSource")
    def test_zhiyun_import_creates_one_resource_per_lecture(self, source_class):
        source = source_class.return_value
        source.list_courses.return_value = [Course(course_id="82493", name="应用统计学", teacher="张老师")]
        source.list_lectures.return_value = [
            Lecture(lecture_id="l1", course_id="82493", title="第一讲", index=1, duration=120),
            Lecture(lecture_id="l2", course_id="82493", title="第二讲", index=2, duration=180),
        ]
        source.get_transcript.side_effect = lambda lecture, refresh=False: [
            TranscriptSegment(lecture_id=lecture.lecture_id, index=0, start_sec=0, end_sec=10, text="课程内容")
        ]
        dataset = self.service.create_dataset(DatasetCreate(name="统计学"))
        resources, warnings = self.service.import_zhiyun_course(dataset.dataset_id, "82493")
        self.assertEqual(warnings, [])
        self.assertEqual([item.title for item in resources], ["第一讲", "第二讲"])
        self.assertEqual(resources[0].current_revision.metadata["lecture_index"], 1)

    @patch("coursebook_agent.product.service.ZhiyunSource")
    def test_zhiyun_import_filters_lectures_and_includes_courseware(self, source_class):
        source = source_class.return_value
        source.list_courses.return_value = [Course(course_id="82493", name="应用统计学")]
        source.list_lectures.return_value = [
            Lecture(lecture_id="l1", course_id="82493", title="第一讲", index=1),
            Lecture(lecture_id="l2", course_id="82493", title="第二讲", index=2),
        ]
        source.get_transcript.return_value = [
            TranscriptSegment(lecture_id="l2", index=0, start_sec=0, end_sec=10, text="课程内容")
        ]
        source.get_courseware.return_value = [
            {"slide_id": "p1", "created_sec": 5, "title": "总体与样本", "image_url": "https://example.test/p1.jpg"}
        ]
        dataset = self.service.create_dataset(DatasetCreate(name="统计学"))
        resources, warnings = self.service.import_zhiyun_course(
            dataset.dataset_id, "82493", lecture_ids=["l2"],
            content_types=["transcript", "courseware"],
        )
        self.assertEqual(warnings, [])
        self.assertEqual([item.kind for item in resources], ["transcript", "courseware"])
        self.assertTrue(all(item.current_revision.metadata["lecture_id"] == "l2" for item in resources))
        self.assertEqual(resources[1].current_revision.page_count, 1)

    @patch("coursebook_agent.product.service.ZhiyunSource")
    def test_zhiyun_import_reports_partial_failures(self, source_class):
        source = source_class.return_value
        source.list_courses.return_value = [Course(course_id="82493", name="应用统计学")]
        source.list_lectures.return_value = [Lecture(lecture_id="l1", course_id="82493", title="第一讲", index=1)]
        source.get_transcript.return_value = [TranscriptSegment(lecture_id="l1", index=0, start_sec=0, end_sec=5, text="内容")]
        source.get_courseware.side_effect = RuntimeError("课件接口不可用")
        dataset = self.service.create_dataset(DatasetCreate(name="统计学"))
        resources, warnings = self.service.import_zhiyun_course(
            dataset.dataset_id, "82493", lecture_ids=["l1"], content_types=["transcript", "courseware"],
        )
        self.assertEqual(len(resources), 1)
        self.assertIn("课件接口不可用", warnings[0])

    @patch("coursebook_agent.product.service.ZhiyunSource")
    def test_zhiyun_import_rejects_unknown_lecture(self, source_class):
        source = source_class.return_value
        source.list_courses.return_value = []
        source.list_lectures.return_value = [Lecture(lecture_id="l1", course_id="82493", title="第一讲", index=1)]
        dataset = self.service.create_dataset(DatasetCreate(name="统计学"))
        with self.assertRaisesRegex(ValueError, "所选讲次不存在"):
            self.service.import_zhiyun_course(dataset.dataset_id, "82493", lecture_ids=["missing"])


class ProductApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.patch = patch("coursebook_agent.product.api.service", return_value=ProductService(Path(self.temp.name)))
        self.patch.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def test_dataset_upload_flow(self):
        created = self.client.post("/api/product/datasets", json={"name": "课程资料"})
        self.assertEqual(created.status_code, 201)
        dataset_id = created.json()["dataset_id"]
        uploaded = self.client.post(
            f"/api/product/datasets/{dataset_id}/resources",
            files={"file": ("outline.txt", "course outline", "text/plain")},
        )
        self.assertEqual(uploaded.status_code, 201)
        revision_id = uploaded.json()["current_revision"]["revision_id"]
        snapshot = self.client.post(
            f"/api/product/datasets/{dataset_id}/snapshots",
            json={"resource_revision_ids": [revision_id], "label": "Run input"},
        )
        self.assertEqual(snapshot.status_code, 201)
        detail = self.client.get(f"/api/product/datasets/{dataset_id}")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(len(detail.json()["resources"]), 1)
        self.assertEqual(len(detail.json()["snapshots"]), 1)

    def test_workflow_preset_is_available(self):
        response = self.client.get("/api/product/workflow-presets")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"][0]["preset_id"], "coursebook")

    def test_run_and_artifact_projections(self):
        from coursebook_agent import app as app_module

        state = JobState(
            job_id="product-run",
            course_id="82493",
            request={"snapshot_id": "snap-1", "lecture_indices": [1, 2]},
            status="running",
            step="生成",
            progress=40,
            message="正在生成章节",
            chapters=[{"index": 1, "title": "第一章", "status": "done"}],
        )
        app_module.jobs[state.job_id] = state
        try:
            run = self.client.get("/api/product/runs/product-run")
            self.assertEqual(run.status_code, 200)
            self.assertEqual(run.json()["snapshot_id"], "snap-1")
            self.assertEqual(run.json()["active_agents"], 1)
            self.assertEqual(run.json()["agents"][0]["status"], "succeeded")

            state.book = CourseBook(course=Course(course_id="82493", name="统计学"), title="统计学教辅")
            artifact = self.client.get("/api/product/artifacts/product-run")
            self.assertEqual(artifact.status_code, 200)
            self.assertEqual(artifact.json()["artifact"]["artifact_id"], "product-run")
        finally:
            app_module.jobs.pop(state.job_id, None)


if __name__ == "__main__":
    unittest.main()
